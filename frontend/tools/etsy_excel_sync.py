#!/usr/bin/env python3
import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


FILE_PATTERN = re.compile(r"^shops_(\d{4}-\d{2}-\d{2})\.xlsx$", re.IGNORECASE)
REQUIRED_HEADERS = {
    "tag": "商品标签",
    "views": "浏览量",
    "favorites": "收藏夹",
    "orders": "订单",
    "revenue": "收入",
}


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", required=True)
    parser.add_argument("--end-date", default="")
    return parser.parse_args()


def parse_money(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group(0)) if match else 0.0


def to_int(value):
    if value is None or value == "":
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def money(value):
    return f"${value:,.2f}"


def number(value):
    return f"{value:,}"


def percent(numerator, denominator):
    if not denominator:
        return "0.0%"
    return f"{(numerator / denominator) * 100:.1f}%"


def parse_date(value):
    return dt.date.fromisoformat(value)


def short_date(value):
    date = parse_date(value)
    return f"{date.month:02d}/{date.day:02d}"


def display_date(value):
    date = parse_date(value)
    return f"{date.month}月{date.day}日"


def stable_id(tag):
    return hashlib.sha1(tag.encode("utf-8")).hexdigest()[:12]


def read_sheet(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError(f"{path.name} 是空表")

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    index = {}
    for key, header in REQUIRED_HEADERS.items():
        if header not in headers:
            raise ValueError(f"{path.name} 缺少列：{header}")
        index[key] = headers.index(header)

    products = []
    for row in rows[1:]:
        tag_value = row[index["tag"]] if index["tag"] < len(row) else ""
        tag = str(tag_value).strip() if tag_value is not None else ""
        if not tag:
            continue
        products.append(
            {
                "id": stable_id(tag),
                "tag": tag,
                "productName": tag,
                "views": to_int(row[index["views"]] if index["views"] < len(row) else 0),
                "favorites": to_int(row[index["favorites"]] if index["favorites"] < len(row) else 0),
                "orders": to_int(row[index["orders"]] if index["orders"] < len(row) else 0),
                "revenue": round(parse_money(row[index["revenue"]] if index["revenue"] < len(row) else 0), 2),
            }
        )
    return products


def load_snapshots(source_dir):
    source = Path(source_dir)
    if not source.exists():
        raise FileNotFoundError(f"找不到 Excel 文件夹：{source}")
    if not source.is_dir():
        raise NotADirectoryError(f"不是文件夹：{source}")

    snapshots = []
    for path in sorted(source.iterdir()):
        if path.name.startswith("~$"):
            continue
        match = FILE_PATTERN.match(path.name)
        if not match:
            continue
        date = match.group(1)
        snapshots.append(
            {
                "date": date,
                "file": path.name,
                "path": str(path),
                "rows": read_sheet(path),
            }
        )
    snapshots.sort(key=lambda item: item["date"])
    return snapshots


def rows_by_tag(rows):
    return {row["tag"]: row for row in rows}


def diff_rows(current_rows, previous_rows=None):
    previous = rows_by_tag(previous_rows or [])
    result = []
    for row in current_rows:
        before = previous.get(row["tag"])
        product = {
            "id": row["id"],
            "tag": row["tag"],
            "productName": row["productName"],
            "views": row["views"] - before["views"] if before else row["views"],
            "favorites": row["favorites"] - before["favorites"] if before else row["favorites"],
            "orders": row["orders"] - before["orders"] if before else row["orders"],
            "revenue": round(row["revenue"] - before["revenue"], 2) if before else row["revenue"],
            "isNew": before is None,
            "note": "Excel 差值同步" if before else "Excel 起始快照",
        }
        result.append(product)
    result.sort(key=lambda item: (item["orders"], item["revenue"], item["views"]), reverse=True)
    return result


def totals(rows):
    return {
        "products": len(rows),
        "views": sum(row["views"] for row in rows),
        "favorites": sum(row["favorites"] for row in rows),
        "orders": sum(row["orders"] for row in rows),
        "revenue": round(sum(row["revenue"] for row in rows), 2),
    }


def best_product(rows):
    if not rows:
        return None
    return sorted(rows, key=lambda item: (item["orders"], item["revenue"], item["favorites"], item["views"]), reverse=True)[0]


def monday_of(date):
    return date - dt.timedelta(days=date.weekday())


def period_start(end_date, period):
    date = parse_date(end_date)
    if period == "day":
        return date
    if period == "week":
        return monday_of(date)
    return date.replace(day=1)


def previous_snapshot_before(snapshots, start_date):
    previous = None
    for snapshot in snapshots:
        if parse_date(snapshot["date"]) < start_date:
            previous = snapshot
    return previous


def snapshots_in_range(snapshots, start_date, end_date):
    return [
        snapshot
        for snapshot in snapshots
        if start_date <= parse_date(snapshot["date"]) <= end_date
    ]


def daily_trends(snapshots, range_snapshots):
    trend_items = []
    for snapshot in range_snapshots:
        index = snapshots.index(snapshot)
        previous = snapshots[index - 1] if index > 0 else None
        rows = diff_rows(snapshot["rows"], previous["rows"] if previous else None)
        total = totals(rows)
        trend_items.append(
            {
                "label": short_date(snapshot["date"]),
                "listings": total["products"],
                "orders": total["orders"],
                "revenue": total["revenue"],
                "adSpend": 0,
                "adRevenue": 0,
                "favorites": total["favorites"],
                "conversations": 0,
            }
        )
    return trend_items


def make_dashboard_period(period, snapshots, end_snapshot):
    end_date = parse_date(end_snapshot["date"])
    start_date = period_start(end_snapshot["date"], period)
    previous = previous_snapshot_before(snapshots, start_date)

    if period == "day":
        end_index = snapshots.index(end_snapshot)
        previous = snapshots[end_index - 1] if end_index > 0 else None
        start_date = end_date
        range_snapshots = snapshots[max(0, end_index - 6): end_index + 1]
    else:
        range_snapshots = snapshots_in_range(snapshots, start_date, end_date)

    product_rows = diff_rows(end_snapshot["rows"], previous["rows"] if previous else None)
    total = totals(product_rows)
    best = best_product(product_rows)
    best_name = best["tag"] if best else "暂无数据"
    best_note = (
        f"{best['orders']} 单 / {money(best['revenue'])} 收入，来自 Excel 同步数据。"
        if best
        else "当前日期没有可用产品数据。"
    )

    period_label = {"day": "按日", "week": "按周", "month": "按月"}[period]
    title_prefix = {"day": display_date(end_snapshot["date"]), "week": "本周", "month": "本月"}[period]
    title = f"{title_prefix}经营总览"
    best_label = {"day": "当日最佳产品", "week": "本周最佳产品", "month": "本月最佳产品"}[period]
    action_title = {"day": "当日建议动作", "week": "本周建议动作", "month": "本月建议动作"}[period]

    if period == "day":
        range_label = end_snapshot["date"]
        trend_title = "当前可用每日趋势"
        source_title = "当日 Excel 数据来源"
    else:
        range_label = f"{start_date.isoformat()} - {end_snapshot['date']}"
        trend_title = f"{title_prefix}每日订单与收入趋势"
        source_title = f"{title_prefix} Excel 数据来源"

    summary = (
        f"{range_label} 已从 Excel 自动同步 {total['products']} 个商品，"
        f"浏览量 {number(total['views'])}，收藏 {number(total['favorites'])}，"
        f"订单 {number(total['orders'])}，收入 {money(total['revenue'])}。"
        f"当前 Excel 只有商品表现字段，广告和站外来源数据等待后续表格接入。"
    )

    metrics = [
        {
            "key": "products",
            "title": "同步商品数",
            "value": number(total["products"]),
            "note": f"来自 {end_snapshot['file']}",
            "tone": "blue",
        },
        {
            "key": "views",
            "title": f"{title_prefix}浏览量",
            "value": number(total["views"]),
            "note": "按 Excel 快照差值计算",
            "tone": "blue",
        },
        {
            "key": "favorites",
            "title": f"{title_prefix}收藏夹",
            "value": number(total["favorites"]),
            "note": f"收藏率 {percent(total['favorites'], total['views'])}",
            "tone": "red",
        },
        {
            "key": "orders",
            "title": f"{title_prefix}订单",
            "value": number(total["orders"]),
            "note": f"转化率 {percent(total['orders'], total['views'])}",
            "tone": "green",
        },
        {
            "key": "revenue",
            "title": f"{title_prefix}收入",
            "value": money(total["revenue"]),
            "note": f"客单价 {money(total['revenue'] / total['orders']) if total['orders'] else '$0.00'}",
            "tone": "green",
        },
        {
            "key": "excel",
            "title": "Excel 同步",
            "value": "已同步",
            "note": f"最新日期 {end_snapshot['date']}",
            "tone": "blue",
        },
        {
            "key": "adSpend",
            "title": "广告数据",
            "value": "待接入",
            "note": "当前表格暂无广告花费列",
            "tone": "amber",
        },
        {
            "key": "traffic",
            "title": "流量来源",
            "value": "待接入",
            "note": "当前表格暂无自然/广告来源列",
            "tone": "amber",
        },
    ]

    return {
        "dashboard": {
            "key": period,
            "label": period_label,
            "eyebrow": "Excel Auto Sync",
            "title": title,
            "rangeLabel": range_label,
            "summary": summary,
            "bestProductLabel": best_label,
            "bestProduct": best_name,
            "bestProductNote": best_note,
            "trendTitle": trend_title,
            "sourceTitle": source_title,
            "actionTitle": action_title,
            "metrics": metrics,
            "trends": daily_trends(snapshots, range_snapshots),
            "trafficSources": [
                {
                    "name": "Excel 商品表现数据",
                    "type": "自然流量",
                    "visits": total["views"],
                    "orders": total["orders"],
                    "revenue": total["revenue"],
                },
                {
                    "name": "广告/站外来源待接入",
                    "type": "广告流量",
                    "visits": 0,
                    "orders": 0,
                    "revenue": 0,
                },
            ],
        },
        "products": product_rows,
    }


def build_response(source_dir, end_date=""):
    snapshots = load_snapshots(source_dir)
    if not snapshots:
        raise FileNotFoundError(f"{source_dir} 下没有 shops_YYYY-MM-DD.xlsx 文件")

    dates = [snapshot["date"] for snapshot in snapshots]
    selected_date = end_date if end_date in dates else dates[-1]
    end_snapshot = next(snapshot for snapshot in snapshots if snapshot["date"] == selected_date)

    periods = {}
    products = {}
    for period in ["day", "week", "month"]:
        result = make_dashboard_period(period, snapshots, end_snapshot)
        periods[period] = result["dashboard"]
        products[period] = result["products"]

    return {
        "ok": True,
        "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
        "sourceDir": source_dir,
        "availableDates": dates,
        "selectedDate": selected_date,
        "latestDate": dates[-1],
        "files": [{"date": item["date"], "name": item["file"], "path": item["path"]} for item in snapshots],
        "periods": periods,
        "products": products,
        "sync": {
            "status": "synced",
            "fileCount": len(snapshots),
            "latestFile": snapshots[-1]["file"],
            "message": f"已读取 {len(snapshots)} 个 Excel 文件",
        },
    }


def main():
    args = parse_args()
    try:
        result = build_response(args.source_dir, args.end_date)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(
            json.dumps(
                {
                    "ok": False,
                    "generatedAt": dt.datetime.now().isoformat(timespec="seconds"),
                    "sourceDir": args.source_dir,
                    "availableDates": [],
                    "selectedDate": "",
                    "latestDate": "",
                    "files": [],
                    "periods": {},
                    "products": {},
                    "sync": {
                        "status": "error",
                        "fileCount": 0,
                        "latestFile": "",
                        "message": str(exc),
                    },
                },
                ensure_ascii=False,
            )
        )
        return 1


if __name__ == "__main__":
    sys.exit(main())
